#!/usr/bin/env python3
"""PPTX builder / editor (python-pptx). All text stays REAL editable text.

Build from job JSON:
  {"size":"16:9"|"4:3", "out":"deck.pptx",
   "slides":[{"bg":"#161616",
              "texts":[{"text":"标题","left":1.2,"top":1.0,"size":44,
                        "color":"#FFFFFF","bold":true,"font":"思源黑体"}],
              "images":[{"image":"a.png","left":8,"top":1,"width":6}]}]}
  Positions in inches. 16:9 = 13.333x7.5in, 4:3 = 10x7.5in.

Edit an existing deck:
  --edit in.pptx --replace-text "旧词=新词" --out out.pptx
"""
import json, sys, argparse
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SIZES = {'16:9': (13.333, 7.5), '4:3': (10, 7.5)}
ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}

def build(job):
    prs = Presentation()
    w, h = SIZES[job.get('size', '16:9')]
    prs.slide_width, prs.slide_height = Inches(w), Inches(h)
    blank = prs.slide_layouts[6]
    for s in job['slides']:
        slide = prs.slides.add_slide(blank)
        if s.get('bg'):
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string(s['bg'].lstrip('#'))
        for t in s.get('texts', []):
            box = slide.shapes.add_textbox(Inches(t.get('left', 0.5)),
                                           Inches(t.get('top', 0.5)),
                                           Inches(t.get('width', w - 1)),
                                           Inches(t.get('height', 1)))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = t['text']
            p.alignment = ALIGN.get(t.get('align', 'left'))
            run = p.runs[0]
            run.font.size = Pt(t.get('size', 24))
            run.font.bold = t.get('bold', False)
            run.font.name = t.get('font', '思源黑体')
            if t.get('color'):
                run.font.color.rgb = RGBColor.from_string(t['color'].lstrip('#'))
        for im in s.get('images', []):
            kw = {}
            if im.get('width'): kw['width'] = Inches(im['width'])
            if im.get('height'): kw['height'] = Inches(im['height'])
            slide.shapes.add_picture(im['image'], Inches(im.get('left', 0)),
                                     Inches(im.get('top', 0)), **kw)
    prs.save(job['out'])
    print('pptx ->', job['out'], len(job['slides']), 'slides')

def edit(args):
    prs = Presentation(args.edit)
    old, new = args.replace_text.split('=', 1)
    n = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if old in r.text:
                            r.text = r.text.replace(old, new)
                            n += 1
    prs.save(args.out or args.edit)
    print(f'replaced {n} run(s) ->', args.out or args.edit)

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('job', nargs='?')
    ap.add_argument('--edit')
    ap.add_argument('--replace-text')
    ap.add_argument('--out')
    a = ap.parse_args()
    if a.edit:
        edit(a)
    else:
        build(json.load(open(a.job)))
